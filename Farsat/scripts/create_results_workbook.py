from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT = Path("results_summary_live_edit_v2.xlsx")

metrics = ["R@10", "R@20", "N@10", "N@20"]

mentor = {
    "Clothing": {
        "R@10": 0.0634,
        "R@20": 0.0936,
        "N@10": 0.0339,
        "N@20": 0.0415,
    },
    "Sports": {
        "R@10": 0.0740,
        "R@20": 0.1122,
        "N@10": 0.0398,
        "N@20": 0.0497,
    },
    "Baby": {
        "R@10": 0.0658,
        "R@20": 0.1005,
        "N@10": 0.0348,
        "N@20": 0.0438,
    },
}

direction_1 = {
    "Baby": {
        "R@10": 0.0737,
        "R@20": 0.1095,
        "N@10": 0.0406,
        "N@20": 0.0498,
    }
}

directions = {
    "Direction 1": direction_1,
    "Direction 2": {},
    "Direction 3": {},
    "Direction 4": {},
    "Direction 5": {},
    "Direction 6": {},
    "Direction 7": {},
}


def style_range(ws):
    thin = Side(style="thin", color="A7B0BC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    sub_fill = PatternFill("solid", fgColor="D9EAF7")

    for row in ws.iter_rows(min_row=1, max_row=5, min_col=1, max_col=75):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        if min_row == 1 and max_row == 1:
            for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
                for cell in row:
                    cell.fill = header_fill
                    cell.font = Font(bold=True, color="FFFFFF", size=13)

    for cell in ws[2]:
        if cell.value is not None:
            cell.fill = sub_fill
            cell.font = Font(bold=True)

    for row in range(3, 6):
        for col in range(1, 76):
            if ws.cell(row, col).value in ("Clothing", "Sports", "Baby"):
                ws.cell(row, col).font = Font(bold=True)

    for col in range(1, 76):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 12
    for col in (1, 7, 17, 27, 37, 47, 57, 67):
        ws.column_dimensions[get_column_letter(col)].width = 14

    ws.freeze_panes = "A3"


def write_mentor_block(ws, start_col):
    end_col = start_col + len(metrics)
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
    ws.cell(1, start_col, "MENTOR")

    headers = ["Dataset", *metrics]
    for idx, header in enumerate(headers, start=start_col):
        ws.cell(2, idx, header)

    for row_idx, dataset in enumerate(["Clothing", "Sports", "Baby"], start=3):
        ws.cell(row_idx, start_col, dataset)
        data = mentor[dataset]
        for metric_idx, metric in enumerate(metrics, start=start_col + 1):
            cell = ws.cell(row_idx, metric_idx, data[metric])
            cell.number_format = "0.0000"


def write_direction_block(ws, start_col, title, rows):
    end_col = start_col + 8
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
    ws.cell(1, start_col, title)

    headers = ["Dataset"]
    for metric in metrics:
        headers.extend([f"{metric} raw", f"{metric} % improvement"])
    for idx, header in enumerate(headers, start=start_col):
        ws.cell(2, idx, header)

    mentor_metric_cols = {"R@10": "B", "R@20": "C", "N@10": "D", "N@20": "E"}
    for row_idx, dataset in enumerate(["Clothing", "Sports", "Baby"], start=3):
        ws.cell(row_idx, start_col, dataset)
        data = rows.get(dataset, {})
        col = start_col + 1
        for metric in metrics:
            raw = data.get(metric)
            raw_cell = ws.cell(row_idx, col, raw)
            raw_ref = f"{get_column_letter(col)}{row_idx}"
            mentor_ref = f"${mentor_metric_cols[metric]}{row_idx}"
            pct_cell = ws.cell(row_idx, col + 1, f'=IF({raw_ref}="","",IFERROR(({raw_ref}-{mentor_ref})/{mentor_ref},""))')
            raw_cell.number_format = "0.0000"
            pct_cell.number_format = "0.00%"
            col += 2


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    write_mentor_block(ws, 1)
    start_col = 7
    for title, rows in directions.items():
        write_direction_block(ws, start_col, title, rows)
        start_col += 10

    style_range(ws)
    wb.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
